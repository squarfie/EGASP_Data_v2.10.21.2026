import re
from collections import Counter
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO, StringIO

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

from .models import Egasp_Data
from django.contrib.auth.decorators import login_required
from django.db.models import Case, When, IntegerField
import csv




CLINICS = ("TLY", "PSH", "TSH")
CLINIC_CODES = {
    "The Love Yourself - Victoria": "TLY",
    "Parañaque Social Hygiene Clinic and Wellness Center": "PSH",
    "Paranaque Social Hygiene Clinic and Wellness Center": "PSH",
    "Taguig Social Hygiene Clinic": "TSH",
}

ARS_HEADERS = """egasp_id clinic consult_date consult_type client_type uic_ptid first_name middle_name last_name suffix birthdate age age_range sex gender_identity gender_identity_other occupation civil_status civil_status_other current_province current_city current_country permanent_province permanent_city permanent_country other_country nationality nationality_other travel_history travel_history_specify client_group client_group_other history_of_sex_partner nationality_sex_partner nationality_sex_partner_other date_of_last_sex number_of_sex_partners relationship_to_partners sb_urethral sb_vaginal sb_anal_insertive sb_oral_insertive sharing_of_sex_toys sb_oral_receptive sb_anal_receptive sb_others sti_none sti_hiv sti_hepatitis_b sti_hepatitis_c sti_ngi sti_syphilis sti_chlamydia sti_anogenital_warts sti_genital_ulcer sti_herpes sti_other sti_trichomoniasis sti_mycoplasma_genitalium sti_lymphogranuloma illicit_drug_use illicit_drug_specify abx_use_prescribed abx_use_prescribed_specify abx_use_self_medicated abx_use_self_medicated_specify abx_use_none abx_use_other abx_use_other_specify route_oral route_injectable_iv route_dermal route_suppository route_other symp_with_discharge symp_no symp_discharge_urethra symp_discharge_vagina symp_discharge_anus symp_discharge_oropharyngeal symp_pain_lower_abdomen symp_tender_testicles symp_painful_urination symp_painful_intercourse symp_rectal_pain symp_other outcome_of_follow_up_visit prev_test_pos prev_test_pos_date result_test_cure_initial result_test_cure_followup no_toc_other_test no_toc_date_performed no_toc_result_of_test patient_compliance_antibiotics other_drugs_specify other_drugs_dosage other_drugs_route other_drugs_duration gonorrhea_treatment treatment_outcome primary_antibiotic primary_abx_other secondary_antibiotic secondary_abx_other notes date_specimen_collection specimen_type specimen_quality date_of_gram_stain diagnosis_at_this_visit gram_neg_intracellular gram_neg_extracellular gs_presence_of_pus_cells presence_gn_intracellular presence_gn_extracellular gs_pus_cells epithelial_cells gs_date_released gs_others gs_negative date_received_in_lab positive_culture_date culture_result growth_span species_identification other_species_id specimen_quality_cs susceptibility_testing_date azm_nd15_val azm_nd15_ris azm_nm_val azm_nm_ris cfm_nd5_val cfm_nd5_ris cfm_nm_val cfm_nm_ris cip_nd5_val cip_nd5_ris cip_nm_val cip_nm_ris cro_nd30_val cro_nd30_ris cro_nm_val cro_nm_ris gen_nd10_val gen_nd10_ris gen_nm_val gen_nm_ris retested_mic confirmation_ast_date beta_lactamase p_png t_rng ars_notes""".split()


EB_HEADERS = """
egasp_id
clinic
consult_date
consult_type
client_type
uic_ptid
birthdate
age
age_range
sex
gender_identity
gender_identity_other
occupation
civil_status
civil_status_other
current_province
current_city
current_country
permanent_province
permanent_city
permanent_country
other_country
nationality
nationality_other
travel_history
travel_history_specify
client_group
client_group_other
history_of_sex_partner
nationality_sex_partner
nationality_sex_partner_other
date_of_last_sex
number_of_sex_partners
relationship_to_partners
sb_urethral
sb_vaginal
sb_anal_insertive
sb_oral_insertive
sharing_of_sex_toys
sb_oral_receptive
sb_anal_receptive
sb_others
sti_none
sti_hiv
sti_hepatitis_b
sti_hepatitis_c
sti_ngi
sti_syphilis
sti_chlamydia
sti_anogenital_warts
sti_genital_ulcer
sti_herpes
sti_other
sti_trichomoniasis
sti_mycoplasma_genitalium
sti_lymphogranuloma
illicit_drug_use
illicit_drug_specify
abx_use_prescribed
abx_use_prescribed_specify
abx_use_self_medicated
abx_use_self_medicated_specify
abx_use_none
abx_use_other
abx_use_other_specify
route_oral
route_injectable_iv
route_dermal
route_suppository
route_other
symp_with_discharge
symp_no
symp_discharge_urethra
symp_discharge_vagina
symp_discharge_anus
symp_discharge_oropharyngeal
symp_pain_lower_abdomen
symp_tender_testicles
symp_painful_urination
symp_painful_intercourse
symp_rectal_pain
symp_other
outcome_of_follow_up_visit
prev_test_pos
prev_test_pos_date
result_test_cure_initial
result_test_cure_followup
no_toc_other_test
no_toc_date_performed
no_toc_result_of_test
patient_compliance_antibiotics
other_drugs_specify
other_drugs_dosage
other_drugs_route
other_drugs_duration
gonorrhea_treatment
treatment_outcome
primary_antibiotic
primary_abx_other
secondary_antibiotic
secondary_abx_other
notes
date_specimen_collection
specimen_type
specimen_quality
date_of_gram_stain
diagnosis_at_this_visit
gram_neg_intracellular
gram_neg_extracellular
gs_presence_of_pus_cells
presence_gn_intracellular
presence_gn_extracellular
gs_pus_cells
epithelial_cells
gs_others
gs_negative
date_received_in_lab
positive_culture_date
culture_result
species_identification
other_species_id
specimen_quality_cs
susceptibility_testing_date
azm_nd15_val
azm_nd15_ris
azm_nm_val
azm_nm_ris
cfm_nd5_val
cfm_nd5_ris
cfm_nm_val
cfm_nm_ris
cip_nd5_val
cip_nd5_ris
cip_nm_val
cip_nm_ris
cro_nd30_val
cro_nd30_ris
cro_nm_val
cro_nm_ris
gen_nd10_val
gen_nd10_ris
gen_nm_val
gen_nm_ris
beta_lactamase
p_png
t_rng
ars_notes
""".split()


ABX_LABELS = {
    "CRO IM 250mg": "Ceftriaxone 250mg",
    "CRO IM 500mg": "Ceftriaxone 500mg",
    "CRO IM 1g": "Ceftriaxone 1g",
    "CFM PO 400mg": "Cefixime 400mg",
    "CFM PO 800mg": "Cefixime 800mg",
    "AZM PO 1g": "Azithromycin 1g",
    "AZM PO 2g": "Azithromycin 2g",
    "DOX PO 100mg": "Doxycycline 100mg",
    "TCY PO 500mg": "Tetracycline 500mg",
}


def clean_header(value):
    value = re.sub(r"[^0-9A-Za-z]+", "_", str(value or "").strip())
    value = re.sub(r"([a-z0-9])([A-Z])", r"\1_\2", value)
    return re.sub(r"_+", "_", value).strip("_").lower()


def clean_value(value):
    if value is None:
        return ""
    if isinstance(value, str):
        value = value.strip()
        return "N/A" if value.lower() == "n/a" else value
    return value


def age_band(value):
    try:
        age = int(float(value))
    except (TypeError, ValueError):
        return ""
    if age < 15:
        return "<15"
    for low, high in ((15, 19), (20, 24), (25, 29), (30, 34), (35, 39), (40, 49), (50, 59)):
        if low <= age <= high:
            return f"{low}-{high}"
    return ">=60"


def _decimal_text(value):
    if value in (None, ""):
        return ""
    if isinstance(value, Decimal):
        return format(value, "f")
    return str(value)


def rows_from_queryset(queryset):
    field_map = {clean_header(field.name): field.name for field in Egasp_Data._meta.fields}
    field_map.update({"p_png": "PPng", "t_rng": "TRng"})
    rows = []
    for record in queryset.prefetch_related("antibiotic_entries").order_by("Egasp_Id"):
        row = {}
        for header in ARS_HEADERS:
            model_name = field_map.get(header)
            row[header] = clean_value(getattr(record, model_name, "")) if model_name else ""
        row["age_range"] = age_band(row.get("age"))
        for entry in record.antibiotic_entries.all():
            code = clean_header(entry.ab_Abx_code)
            if code and f"{code}_val" in row:
                if entry.ab_Disk_value is not None:
                    row[f"{code}_val"] = entry.ab_Disk_value
                    row[f"{code}_ris"] = entry.ab_Disk_RIS or ""
                elif entry.ab_MIC_value is not None:
                    row[f"{code}_val"] = f"{entry.ab_MIC_operand or ''}{_decimal_text(entry.ab_MIC_value)}"
                    row[f"{code}_ris"] = entry.ab_MIC_RIS or ""
        rows.append(row)
    return rows

############## EB STYLE REPORT queryset #######################

def rows_from_queryset_eb(queryset):
    field_map = {
        clean_header(field.name): field.name
        for field in Egasp_Data._meta.fields
    }

    field_map.update({
        "p_png": "PPng",
        "t_rng": "TRng",
    })

    rows = []

    for record in queryset.prefetch_related(
        "antibiotic_entries"
    ).order_by("Egasp_Id"):

        row = {}

        for header in EB_HEADERS:
            model_name = field_map.get(header)

            row[header] = (
                clean_value(
                    getattr(record, model_name, "")
                )
                if model_name
                else ""
            )

        # Antibiotic values
        for entry in record.antibiotic_entries.all():

            code = clean_header(entry.ab_Abx_code)

            value_field = f"{code}_val"
            ris_field = f"{code}_ris"

            if value_field not in EB_HEADERS:
                continue

            if entry.ab_Disk_value is not None:
                row[value_field] = entry.ab_Disk_value

                if ris_field in EB_HEADERS:
                    row[ris_field] = (
                        entry.ab_Disk_RIS or ""
                    )

            elif entry.ab_MIC_value is not None:
                row[value_field] = (
                    f"{entry.ab_MIC_operand or ''}"
                    f"{_decimal_text(entry.ab_MIC_value)}"
                )

                if ris_field in EB_HEADERS:
                    row[ris_field] = (
                        entry.ab_MIC_RIS or ""
                    )

        rows.append(row)

    return rows

######################## end of EB report ###########################



# def rows_from_upload(uploaded_file):
#     workbook = load_workbook(uploaded_file, read_only=True, data_only=True)
#     sheet = workbook["ARS Data"] if "ARS Data" in workbook.sheetnames else workbook.active
#     values = sheet.iter_rows(values_only=True)
#     try:
#         headers = [clean_header(value) for value in next(values)]
#     except StopIteration as exc:
#         raise ValueError("The uploaded workbook is empty.") from exc
#     if "egasp_id" not in headers or "clinic" not in headers:
#         raise ValueError("The uploaded list must include egasp_id and clinic columns.")
#     rows = []
#     for values_row in values:
#         if not any(value not in (None, "") for value in values_row):
#             continue
#         row = {header: clean_value(value) for header, value in zip(headers, values_row) if header}
#         row["age_range"] = row.get("age_range") or age_band(row.get("age"))
#         rows.append(row)
#     return rows





def rows_from_upload(uploaded_file):
    filename = uploaded_file.name.lower()

    # =========================================================
    # XLSX
    # =========================================================
    if filename.endswith(".xlsx"):

        workbook = load_workbook(
            uploaded_file,
            read_only=True,
            data_only=True,
        )

        sheet = (
            workbook["ARS Data"]
            if "ARS Data" in workbook.sheetnames
            else workbook.active
        )

        values = sheet.iter_rows(values_only=True)

        try:
            headers = [
                clean_header(value)
                for value in next(values)
            ]
        except StopIteration as exc:
            raise ValueError(
                "The uploaded workbook is empty."
            ) from exc

        data_rows = values

    # =========================================================
    # CSV / TSV
    # =========================================================
    elif filename.endswith((".csv", ".tsv")):

        try:
            uploaded_file.seek(0)

            text = uploaded_file.read().decode(
                "utf-8-sig"
            )

        except UnicodeDecodeError as exc:
            raise ValueError(
                "The uploaded CSV/TSV file must use UTF-8 encoding."
            ) from exc

        delimiter = (
            "\t"
            if filename.endswith(".tsv")
            else ","
        )

        reader = csv.reader(
            StringIO(text),
            delimiter=delimiter,
        )

        try:
            headers = [
                clean_header(value)
                for value in next(reader)
            ]
        except StopIteration as exc:
            raise ValueError(
                "The uploaded file is empty."
            ) from exc

        data_rows = reader

    # =========================================================
    # UNSUPPORTED
    # =========================================================
    else:
        raise ValueError(
            "Unsupported file type. "
            "Upload an XLSX, CSV, or TSV file."
        )

    # =========================================================
    # REQUIRED COLUMNS
    # =========================================================
    if (
        "egasp_id" not in headers
        or "clinic" not in headers
    ):
        raise ValueError(
            "The uploaded list must include "
            "egasp_id and clinic columns."
        )

    # =========================================================
    # BUILD ROW DICTIONARIES
    # =========================================================
    rows = []

    for values_row in data_rows:

        if not any(
            value not in (None, "")
            for value in values_row
        ):
            continue

        row = {
            header: clean_value(value)
            for header, value in zip(
                headers,
                values_row,
            )
            if header
        }

        row["age_range"] = (
            row.get("age_range")
            or age_band(row.get("age"))
        )

        rows.append(row)

    return rows


def _clinic(row):
    return CLINIC_CODES.get(str(row.get("clinic", "")).strip())


def _percent(number, total):
    return round(number / total * 100, 1) if total else 0


def _summary_row(category, label, counts, denominators=None):
    denominators = denominators or counts
    numbers = [counts.get(clinic, 0) for clinic in CLINICS]
    percents = [_percent(counts.get(clinic, 0), denominators.get(clinic, 0)) for clinic in CLINICS]
    return [category, label, *numbers, *percents]


def _category_rows(rows, field, category):
    counters = {clinic: Counter() for clinic in CLINICS}
    for row in rows:
        clinic = _clinic(row)
        value = clean_value(row.get(field))
        if clinic and value not in (None, ""):
            if field in ("primary_antibiotic", "secondary_antibiotic"):
                value = ABX_LABELS.get(value, value)
            counters[clinic][str(value)] += 1
    denominators = {clinic: sum(counter.values()) for clinic, counter in counters.items()}
    labels = {label for counter in counters.values() for label in counter}
    if field == "age_range":
        age_order = ("<15", "15-19", "20-24", "25-29", "30-34", "35-39", "40-49", "50-59", ">=60")
        labels = [label for label in age_order if label in labels]
    else:
        labels = sorted(labels)
    return [
        _summary_row(
            category,
            label,
            {clinic: counters[clinic][label] for clinic in CLINICS},
            denominators,
        )
        for label in labels
    ]





def build_summary(rows):
    output = []
    patient_ids = {clinic: set() for clinic in CLINICS}
    for index, row in enumerate(rows):
        clinic = _clinic(row)
        if clinic:
            patient_ids[clinic].add(row.get("egasp_id") or f"__row_{index}")
    total_counts = Counter(
        {clinic: len(values) for clinic, values in patient_ids.items()}
    )
    total_patients = sum(total_counts.values())
    output.append(
        _summary_row(
            "Samples Summary",
            "Total number of patients",
            total_counts,
            {clinic: total_patients for clinic in CLINICS},
        )
    )
    for label, predicate in (
        ("Gram Stain Positive specimen", lambda r: r.get("presence_gn_intracellular") == "Yes" or r.get("presence_gn_extracellular") == "Yes"),
        ("Culture Positive for Neisseria gonorrhoeae", lambda r: r.get("culture_result") == "Positive"),
    ):
        counts = Counter(_clinic(row) for row in rows if _clinic(row) and predicate(row))
        total = sum(counts.values())
        denominators = {clinic: total for clinic in CLINICS}
        output.append(_summary_row("Samples Summary", label, counts, denominators))
    output.extend(_category_rows(rows, "primary_antibiotic", "Primary antibiotic"))
    output.extend(_category_rows(rows, "secondary_antibiotic", "Secondary antibiotic"))
    output.extend(_category_rows(rows, "age_range", "Age range"))
    output.extend(_category_rows(rows, "history_of_sex_partner", "History of sex partner"))
    abx_rows = [dict(row, abx_use="Yes" if row.get("abx_use_none") == "No" else "No") for row in rows]
    output.extend(_category_rows(abx_rows, "abx_use", "Antibiotic use in last 2 weeks"))
    return output


def build_mic_summary(rows):
    output = []
    for field, category in (
        ("azm_nm_val", "Azithromycin MIC"),
        ("cfm_nm_val", "Cefixime MIC"),
        ("cro_nm_val", "Ceftriaxone MIC"),
    ):
        output.extend(_category_rows(rows, field, category))
    return output



# def build_report_workbook(rows):
#     workbook = Workbook()
#     workbook.remove(workbook.active)
#     sheets = (
#         ("ARS Data", ARS_HEADERS, [[row.get(header, "") for header in ARS_HEADERS] for row in rows]),
#         ("Summary", ["Category", "Row", "TLY_No", "PSH_No", "TSH_No", "TLY_%", "PSH_%", "TSH_%"], build_summary(rows)),
#         ("MIC Summary", ["Category", "Row", "TLY_No", "PSH_No", "TSH_No", "TLY_%", "PSH_%", "TSH_%"], build_mic_summary(rows)),
#     )
#     for table_number, (title, headers, data_rows) in enumerate(sheets, start=1):
#         sheet = workbook.create_sheet(title)
#         sheet.append(headers)
#         for data_row in data_rows:
#             sheet.append(data_row)
#         table = Table(displayName=f"EGASPReportTable{table_number}", ref=sheet.dimensions)
#         table.tableStyleInfo = TableStyleInfo(
#             name="TableStyleLight9",
#             showFirstColumn=False,
#             showLastColumn=False,
#             showRowStripes=True,
#             showColumnStripes=False,
#         )
#         sheet.add_table(table)
#         for column in range(1, sheet.max_column + 1):
#             values = [str(sheet.cell(row, column).value or "") for row in range(1, min(sheet.max_row, 100) + 1)]
#             sheet.column_dimensions[get_column_letter(column)].width = min(max(map(len, values), default=8) + 2, 42)
#     output = BytesIO()
#     workbook.save(output)
#     return output.getvalue()


def build_report_workbook(rows):
    workbook = Workbook()
    workbook.remove(workbook.active)

    sheets = (
        (
            "ARS Data",
            ARS_HEADERS,
            [
                [row.get(header, "") for header in ARS_HEADERS]
                for row in rows
            ],
        ),
        (
            "Summary",
            [
                "Category",
                "Row",
                "TLY_No",
                "PSH_No",
                "TSH_No",
                "TLY_%",
                "PSH_%",
                "TSH_%",
            ],
            build_summary(rows),
        ),
        (
            "MIC Summary",
            [
                "Category",
                "Row",
                "TLY_No",
                "PSH_No",
                "TSH_No",
                "TLY_%",
                "PSH_%",
                "TSH_%",
            ],
            build_mic_summary(rows),
        ),
    )

    for table_number, (title, headers, data_rows) in enumerate(
        sheets,
        start=1,
    ):
        sheet = workbook.create_sheet(title)

        # Header row
        sheet.append(headers)

        # Data rows
        for data_row in data_rows:
            sheet.append(data_row)

        # --------------------------------------------------
        # IMPORTANT:
        # Only create an Excel Table when there is actual data
        # --------------------------------------------------
        if data_rows:
            last_column = get_column_letter(len(headers))
            last_row = sheet.max_row

            table_ref = f"A1:{last_column}{last_row}"

            table = Table(
                displayName=f"EGASPReportTable{table_number}",
                ref=table_ref,
            )

            table.tableStyleInfo = TableStyleInfo(
                name="TableStyleLight9",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )

            sheet.add_table(table)

        # Optional: keep header visible
        sheet.freeze_panes = "A2"

        # Column widths
        for column in range(1, sheet.max_column + 1):
            values = [
                str(sheet.cell(row, column).value or "")
                for row in range(
                    1,
                    min(sheet.max_row, 100) + 1
                )
            ]

            sheet.column_dimensions[
                get_column_letter(column)
            ].width = min(
                max(map(len, values), default=8) + 2,
                42,
            )

    output = BytesIO()
    workbook.save(output)

    return output.getvalue()



#################### EB STYLE WORKBOOK ################################

def build_eb_report_workbook(rows):
    workbook = Workbook()
    workbook.remove(workbook.active)

    # Create EB Data sheet only
    sheet = workbook.create_sheet("EB Data")

    # Headers
    sheet.append(EB_HEADERS)

    # Data
    for row in rows:
        sheet.append([
            row.get(header, "")
            for header in EB_HEADERS
        ])

    # Add Excel table only if there is actual data
    if rows:
        last_column = get_column_letter(len(EB_HEADERS))
        last_row = sheet.max_row

        table_ref = f"A1:{last_column}{last_row}"

        table = Table(
            displayName="EGASPEBReportTable",
            ref=table_ref,
        )

        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleLight9",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )

        sheet.add_table(table)

    # Freeze header
    sheet.freeze_panes = "A2"

    # Auto-size columns
    for column in range(1, sheet.max_column + 1):
        values = [
            str(sheet.cell(row, column).value or "")
            for row in range(
                1,
                min(sheet.max_row, 100) + 1
            )
        ]

        sheet.column_dimensions[
            get_column_letter(column)
        ].width = min(
            max(map(len, values), default=8) + 2,
            42,
        )

    # Save workbook to memory
    output = BytesIO()
    workbook.save(output)

    return output.getvalue()

####################### end of eb workbook ##################




################# Match egasp_id from database ##############


def _parse_egasp_ids(raw_ids):
    """
    Accepts:
        RL260166, RL260167, RL260168

    or values pasted directly from an Excel column:
        RL260166
        RL260167
        RL260168
    """

    values = re.split(r"[,\s]+", raw_ids or "")

    egasp_ids = []

    for value in values:
        egasp_id = value.strip().upper()

        if egasp_id and egasp_id not in egasp_ids:
            egasp_ids.append(egasp_id)

    return egasp_ids



def _get_egasp_queryset_in_order(egasp_ids):
    ordering = Case(
        *[
            When(
                Egasp_Id=egasp_id,
                then=position,
            )
            for position, egasp_id in enumerate(egasp_ids)
        ],
        output_field=IntegerField(),
    )

    queryset = (
        Egasp_Data.objects
        .filter(Egasp_Id__in=egasp_ids)
        .order_by(ordering)
    )

    found_ids = set(
        queryset.values_list("Egasp_Id", flat=True)
    )

    missing_ids = [
        egasp_id
        for egasp_id in egasp_ids
        if egasp_id not in found_ids
    ]

    return queryset, missing_ids


