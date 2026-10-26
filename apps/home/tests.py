# -*- encoding: utf-8 -*-
"""
Copyright (c) 2019 - present AppSeed.us
"""

import csv
from datetime import date
from io import BytesIO, StringIO

from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.template.loader import render_to_string
from django.test import RequestFactory
from django.urls import resolve, reverse
from openpyxl import Workbook, load_workbook

from .models import AntibioticEntry, Egasp_Data
from .utils import clean_date

# Create your tests here.


class CleanDateTests(TestCase):
    def test_clean_date_accepts_month_day_year_import_format(self):
        self.assertEqual(clean_date("10/23/2026"), date(2026, 10, 23))


class SidenavActiveStateTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_superuser(
            username="navigation-admin", password="password"
        )
        self.request_factory = RequestFactory()

    def test_each_sidenav_page_link_is_highlighted_on_its_page(self):
        route_names = (
            "home",
            "crf_data",
            "show_data",
            "monthly_report",
            "add_breakpoints",
            "add_dropdown",
            "add_specimen",
            "add_contact",
            "add_clin_pers",
            "add_location",
        )

        for route_name in route_names:
            with self.subTest(route_name=route_name):
                url = reverse(route_name)
                request = self.request_factory.get(url)
                request.user = self.user
                request.resolver_match = resolve(url)
                rendered = render_to_string(
                    "includes/sidenav.html", request=request
                )

                self.assertIn(
                    f'class="nav-link active" href="{url}"', rendered
                )


class MonthlyReportTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(
            username="report-user", password="password"
        )
        self.client.force_login(self.user)
        self.url = reverse("monthly_report")

    def _workbook_from_response(self, response):
        return load_workbook(BytesIO(response.content), data_only=True)

    def test_report_page_is_available_to_authenticated_users(self):
        response = self.client.get(self.url)

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'name="source" value="upload"')
        self.assertContains(response, 'name="source" value="database"')
        self.assertContains(response, 'name="from_date"')
        self.assertContains(response, 'name="to_date"')
        self.assertContains(
            response,
            'type="file" class="form-control-file" id="report-file"',
        )
        self.assertNotContains(response, 'class="custom-file-label"')

    def test_database_report_filters_inclusively_by_consult_date(self):
        tly_record = Egasp_Data.objects.create(
            Egasp_Id="RL260001",
            Clinic="The Love Yourself - Victoria",
            Consult_Date=date(2026, 9, 1),
            Age="21",
            Culture_Result="Positive",
        )
        AntibioticEntry.objects.create(
            ab_idNumber_egasp=tly_record,
            ab_EgaspId=tly_record.Egasp_Id,
            ab_Antibiotic="Azithromycin",
            ab_Abx_code="AZM_NM",
            ab_MIC_value="0.064",
            ab_MIC_RIS="S",
        )
        Egasp_Data.objects.create(
            Egasp_Id="RP260001",
            Clinic="Parañaque Social Hygiene Clinic and Wellness Center",
            Consult_Date=date(2026, 9, 30),
            Age="30",
            Culture_Result="Positive",
        )
        Egasp_Data.objects.create(
            Egasp_Id="RT260001",
            Clinic="Taguig Social Hygiene Clinic",
            Consult_Date=date(2026, 10, 1),
            Age="40",
        )

        response = self.client.post(
            self.url,
            {
                "source": "database",
                "from_date": "2026-09-01",
                "to_date": "2026-09-30",
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        workbook = self._workbook_from_response(response)
        self.assertEqual(workbook.sheetnames, ["ARS Data", "Summary", "MIC Summary"])
        ids = [cell.value for cell in workbook["ARS Data"]["A"][1:]]
        self.assertEqual(ids, ["RL260001", "RP260001"])
        summary_rows = list(workbook["Summary"].iter_rows(values_only=True))
        total_row = next(row for row in summary_rows if row[1] == "Total number of patients")
        self.assertEqual(total_row[2:5], (1, 1, 0))
        mic_rows = list(workbook["MIC Summary"].iter_rows(values_only=True))
        self.assertIn(
            ("Azithromycin MIC", "0.064", 1, 0, 0, 100, 0, 0),
            mic_rows,
        )

    def test_uploaded_ars_list_generates_summary_sheets(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "ARS Data"
        sheet.append([
            "egasp_id", "clinic", "consult_date", "age",
            "primary_antibiotic", "secondary_antibiotic",
            "history_of_sex_partner", "abx_use_none",
            "presence_gn_intracellular", "presence_gn_extracellular",
            "culture_result", "azm_nm_val", "cfm_nm_val", "cro_nm_val",
        ])
        sheet.append([
            "RL260001", "The Love Yourself - Victoria", date(2026, 9, 1), 21,
            "CRO IM 1g", "DOX PO 100mg", "with Male", "No",
            "Yes", "No", "Positive", "0.064", "", "<0.016",
        ])
        payload = BytesIO()
        workbook.save(payload)
        upload = SimpleUploadedFile(
            "egasp-list.xlsx",
            payload.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

        response = self.client.post(
            self.url, {"source": "upload", "report_file": upload}
        )

        self.assertEqual(response.status_code, 200)
        result = self._workbook_from_response(response)
        summary_rows = list(result["Summary"].iter_rows(values_only=True))
        self.assertIn(
            ("Primary antibiotic", "Ceftriaxone 1g", 1, 0, 0, 100, 0, 0),
            summary_rows,
        )
        mic_rows = list(result["MIC Summary"].iter_rows(values_only=True))
        self.assertIn(
            ("Azithromycin MIC", "0.064", 1, 0, 0, 100, 0, 0),
            mic_rows,
        )

    def test_database_report_rejects_reversed_date_range(self):
        response = self.client.post(
            self.url,
            {
                "source": "database",
                "from_date": "2026-10-01",
                "to_date": "2026-09-01",
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "The start date must be on or before the end date.")


class AntibioticUploadTests(TestCase):
    headers = [
        "id",
        "ab_EgaspId",
        "ab_Antibiotic",
        "ab_Abx_code",
        "ab_Abx",
        "ab_Disk_value",
        "ab_Disk_RIS",
        "ab_MIC_operand",
        "ab_MIC_value",
        "ab_MIC_RIS",
        "ab_AlertMIC",
        "ab_Alert_val",
        "ab_Retest_Antibiotic",
        "ab_Retest_Abx_code",
        "ab_Retest_Abx",
        "ab_Retest_DiskValue",
        "ab_Retest_Disk_RIS",
        "ab_Retest_MIC_operand",
        "ab_Retest_MICValue",
        "ab_Retest_MIC_RIS",
        "ab_Retest_AlertMIC",
        "ab_Retest_Alert_val",
        "ab_R_breakpoint",
        "ab_I_breakpoint",
        "ab_SDD_breakpoint",
    ]

    def setUp(self):
        self.admin = User.objects.create_user(
            username="upload-admin", password="password", is_staff=True
        )
        self.client.force_login(self.admin)
        self.egasp = Egasp_Data.objects.create(Egasp_Id="RL180002")
        self.url = reverse("upload_antibiotic_entries")

    def sample_row(self, **overrides):
        row = {
            "id": "96",
            "ab_EgaspId": "RL180002",
            "ab_Antibiotic": "Azithromycin",
            "ab_Abx_code": "AZM_NM",
            "ab_Abx": "AZM",
            "ab_Disk_value": "",
            "ab_Disk_RIS": "",
            "ab_MIC_operand": "",
            "ab_MIC_value": "0.06",
            "ab_MIC_RIS": "S",
            "ab_AlertMIC": "FALSE",
            "ab_Alert_val": "",
            "ab_Retest_Antibiotic": "Azithromycin",
            "ab_Retest_Abx_code": "AZM_NM",
            "ab_Retest_Abx": "AZM",
            "ab_Retest_DiskValue": "",
            "ab_Retest_Disk_RIS": "",
            "ab_Retest_MIC_operand": "<=",
            "ab_Retest_MICValue": "0.03",
            "ab_Retest_MIC_RIS": "S",
            "ab_Retest_AlertMIC": "FALSE",
            "ab_Retest_Alert_val": "",
            "ab_R_breakpoint": "",
            "ab_I_breakpoint": "",
            "ab_SDD_breakpoint": "",
        }
        row.update(overrides)
        return row

    def csv_upload(self, rows, headers=None, name="antibiotic_entries_ready.csv"):
        output = StringIO()
        writer = csv.DictWriter(output, fieldnames=headers or self.headers)
        writer.writeheader()
        writer.writerows(rows)
        return SimpleUploadedFile(
            name, output.getvalue().encode("utf-8"), content_type="text/csv"
        )

    def xlsx_upload(self, rows):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(self.headers)
        for row in rows:
            sheet.append([row.get(header, "") for header in self.headers])
        output = BytesIO()
        workbook.save(output)
        return SimpleUploadedFile(
            "antibiotic_entries_ready.xlsx",
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    def test_csv_upload_connects_antibiotic_to_matching_egasp_id(self):
        response = self.client.post(
            self.url,
            {"File_uploadAbx": self.csv_upload([self.sample_row()])},
            follow=True,
        )

        entry = AntibioticEntry.objects.get()
        self.assertEqual(entry.ab_idNumber_egasp, self.egasp)
        self.assertEqual(entry.ab_EgaspId, "RL180002")
        self.assertEqual(entry.ab_Abx_code, "AZM_NM")
        self.assertEqual(str(entry.ab_MIC_value), "0.060")
        self.assertEqual(str(entry.ab_Retest_MICValue), "0.030")
        self.assertContains(response, "1 created, 0 updated")

    def test_reupload_updates_existing_antibiotic_instead_of_duplicating_it(self):
        first = self.csv_upload([self.sample_row()])
        second = self.csv_upload([self.sample_row(ab_MIC_value="0.12")])

        self.client.post(self.url, {"File_uploadAbx": first})
        response = self.client.post(
            self.url, {"File_uploadAbx": second}, follow=True
        )

        self.assertEqual(AntibioticEntry.objects.count(), 1)
        self.assertEqual(str(AntibioticEntry.objects.get().ab_MIC_value), "0.120")
        self.assertContains(response, "0 created, 1 updated")

    def test_xlsx_upload_is_supported_and_connected(self):
        response = self.client.post(
            self.url,
            {"File_uploadAbx": self.xlsx_upload([self.sample_row()])},
            follow=True,
        )

        self.assertEqual(AntibioticEntry.objects.get().ab_idNumber_egasp, self.egasp)
        self.assertContains(response, "1 created, 0 updated")

    def test_missing_egasp_ids_are_reported_to_the_user(self):
        upload = self.csv_upload([self.sample_row(ab_EgaspId="RL999999")])

        response = self.client.post(
            self.url, {"File_uploadAbx": upload}, follow=True
        )

        self.assertFalse(AntibioticEntry.objects.exists())
        self.assertContains(response, "Missing EGASP IDs: RL999999")

    def test_missing_required_headers_rejects_file_without_saving_rows(self):
        headers = ["ab_Antibiotic", "ab_MIC_value"]
        upload = self.csv_upload(
            [{"ab_Antibiotic": "Azithromycin", "ab_MIC_value": "0.06"}],
            headers=headers,
        )

        response = self.client.post(
            self.url, {"File_uploadAbx": upload}, follow=True
        )

        self.assertFalse(AntibioticEntry.objects.exists())
        self.assertContains(
            response,
            "Missing required columns: ab_EgaspId, ab_Abx_code",
        )


class DashboardPaginationTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(
            username="dashboard-user", password="password"
        )
        self.client.force_login(self.user)
        for number in range(1, 31):
            Egasp_Data.objects.create(
                Egasp_Id=f"RL{number:06d}",
                First_Name=f"Patient {number}",
                Last_Name="Example",
                Clinic="Test Clinic",
            )

    def test_dashboard_defaults_to_ten_records_per_page(self):
        response = self.client.get(reverse("home"))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.context["isolates"]), 10)
        self.assertEqual(response.context["isolates"].paginator.count, 30)
        self.assertContains(response, "Showing 1 to 10 of 30 entries")
        self.assertContains(response, "Page 1 of 3")

    def test_dashboard_can_show_requested_page_and_page_size(self):
        response = self.client.get(
            reverse("home"), {"page": 2, "page_size": 25}
        )

        self.assertEqual(len(response.context["isolates"]), 5)
        self.assertEqual(response.context["isolates"].number, 2)
        self.assertEqual(response.context["page_size"], 25)
        self.assertContains(response, "Showing 26 to 30 of 30 entries")

    def test_dashboard_searches_all_records_before_paginating(self):
        response = self.client.get(reverse("home"), {"q": "RL000021"})

        self.assertEqual(response.context["isolates"].paginator.count, 1)
        self.assertContains(response, "RL000021")
        self.assertNotContains(response, "RL000020")

    def test_dashboard_pagination_links_preserve_search_and_page_size(self):
        response = self.client.get(
            reverse("home"), {"q": "Patient", "page_size": 10}
        )

        self.assertContains(response, "?q=Patient&amp;page_size=10&amp;page=2")

    def test_dashboard_rejects_unsupported_page_size(self):
        response = self.client.get(reverse("home"), {"page_size": 9999})

        self.assertEqual(response.context["page_size"], 10)
        self.assertEqual(len(response.context["isolates"]), 10)

    def test_dashboard_keeps_legacy_argon_layout_with_server_side_controls(self):
        response = self.client.get(reverse("home"))

        self.assertContains(response, 'class="card card-stats"')
        self.assertContains(response, 'class="table align-items-center table-flush table-hover"')
        self.assertContains(response, 'name="q"')
        self.assertContains(response, 'name="page_size"')
        self.assertNotContains(response, 'class="dashboard-stat')
